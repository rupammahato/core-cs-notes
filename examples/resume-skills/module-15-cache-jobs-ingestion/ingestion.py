from dataclasses import dataclass
from datetime import datetime,timezone
from hashlib import sha256
from io import StringIO,BytesIO
import csv,re,zipfile
from defusedxml.ElementTree import fromstring
from openpyxl import load_workbook
from pypdf import PdfReader
from store import transaction

FIELDS=['external_id','name','amount_minor','source_version','source_time']
MAX_BYTES=2_000_000
MAX_ROWS=1000
csv.field_size_limit(2048)

class InvalidFile(ValueError):pass

@dataclass(frozen=True)
class Record:
    external_id:str
    name:str
    amount_minor:int
    source_version:int
    source_time:str

def bounded(data):
    if not isinstance(data,bytes) or len(data)>MAX_BYTES:raise InvalidFile('file byte limit')
    return data

def normalize(rows):
    records=[];seen=set()
    for number,row in enumerate(rows,1):
        if number>MAX_ROWS:raise InvalidFile('row limit')
        if len(row)!=len(FIELDS):raise InvalidFile(f'record {number}: column count')
        values=[str(v).strip() if v is not None else '' for v in row]
        identifier,name,amount,version,stamp=values
        if not re.fullmatch(r'[A-Za-z0-9_-]{1,64}',identifier) or not 1<=len(name)<=80:
            raise InvalidFile(f'record {number}: identity/name')
        if identifier in seen:raise InvalidFile(f'record {number}: duplicate source identity')
        if not re.fullmatch(r'\d{1,12}',amount) or not re.fullmatch(r'\d{1,9}',version):
            raise InvalidFile(f'record {number}: integer amount/version')
        try:
            parsed=datetime.fromisoformat(stamp)
            if parsed.tzinfo is None:raise ValueError('timezone required')
        except ValueError as error:raise InvalidFile(f'record {number}: aware timestamp required') from error
        records.append(Record(identifier,name,int(amount),int(version),parsed.astimezone(timezone.utc).isoformat()))
        seen.add(identifier)
    if not records:raise InvalidFile('empty data')
    return records

def parse_csv(data):
    try:
        reader=csv.reader(StringIO(bounded(data).decode('utf-8-sig'),newline=''),strict=True)
        if next(reader,None)!=FIELDS:raise InvalidFile('exact CSV header required')
        return normalize(reader)
    except (UnicodeError,csv.Error) as error:raise InvalidFile('CSV encoding/structure') from error

def parse_xml(data):
    root=fromstring(bounded(data),forbid_dtd=True,forbid_entities=True,forbid_external=True)
    if root.tag!='records':raise InvalidFile('XML root')
    rows=[]
    for element in root:
        if element.tag!='record' or [child.tag for child in element]!=FIELDS:raise InvalidFile('XML record shape')
        rows.append([child.text for child in element])
    return normalize(rows)

def parse_xlsx(data):
    stream=BytesIO(bounded(data))
    with zipfile.ZipFile(stream) as archive:
        info=archive.infolist()
        if len(info)>200 or sum(x.file_size for x in info)>8_000_000:raise InvalidFile('XLSX archive budget')
        if any('vbaproject' in x.filename.lower() for x in info):raise InvalidFile('macro content rejected')
    stream.seek(0)
    workbook=load_workbook(stream,read_only=True,data_only=False,keep_links=False)
    try:
        if workbook.sheetnames!=['Records']:raise InvalidFile('one Records worksheet required')
        sheet=workbook['Records'];sheet.reset_dimensions()
        rows=sheet.iter_rows();header=next(rows,())
        if [cell.value for cell in header]!=FIELDS:raise InvalidFile('XLSX header')
        def values():
            for row in rows:
                if any(cell.data_type=='f' for cell in row):raise InvalidFile('formula cells rejected')
                yield [cell.value for cell in row]
        return normalize(values())
    finally:workbook.close()

def extract_pdf(data,*,max_pages=5):
    reader=PdfReader(BytesIO(bounded(data)),strict=True)
    if reader.is_encrypted:raise InvalidFile('encrypted PDF unsupported')
    if len(reader.pages)>max_pages:raise InvalidFile('PDF page limit')
    text='\n'.join(page.extract_text() or '' for page in reader.pages)
    if not text.strip():raise InvalidFile('no extractable text; OCR/review required')
    if len(text)>100_000:raise InvalidFile('PDF extracted-text budget')
    return text.strip()  # Evidence text, not structured financial records.

def import_csv(conn,*,tenant_id,import_key,data,batch_size=2,fail_batch=None,after_commit=None):
    if type(tenant_id) is not int or tenant_id<1 or type(import_key) is not str or not 1<=len(import_key)<=80:
        raise ValueError('invalid trusted import scope')
    if type(batch_size) is not int or not 1<=batch_size<=100:raise ValueError('invalid batch size')
    records=parse_csv(data);digest=sha256(data).hexdigest()
    with transaction(conn):
        conn.execute('INSERT INTO imports(tenant_id,import_key,digest) VALUES(?,?,?) ON CONFLICT(tenant_id,import_key) DO NOTHING',(tenant_id,import_key,digest))
        job=conn.execute('SELECT * FROM imports WHERE tenant_id=? AND import_key=?',(tenant_id,import_key)).fetchone()
        if job['digest']!=digest:raise ValueError('import key reused with changed bytes')
    while True:
        with transaction(conn):
            job=conn.execute('SELECT * FROM imports WHERE tenant_id=? AND import_key=?',(tenant_id,import_key)).fetchone()
            start=job['checkpoint']
            if start==len(records):return start
            end=min(start+batch_size,len(records))
            for row in records[start:end]:
                previous=conn.execute('SELECT * FROM records WHERE tenant_id=? AND external_id=?',(tenant_id,row.external_id)).fetchone()
                if previous and previous['source_version']==row.source_version:
                    if (previous['name'],previous['amount_minor'],previous['source_time'])!=(row.name,row.amount_minor,row.source_time):
                        raise ValueError('same source version has conflicting content')
                conn.execute('''INSERT INTO records VALUES(?,?,?,?,?,?) ON CONFLICT(tenant_id,external_id) DO UPDATE SET
                    name=excluded.name,amount_minor=excluded.amount_minor,source_version=excluded.source_version,source_time=excluded.source_time
                    WHERE excluded.source_version>records.source_version''',
                    (tenant_id,row.external_id,row.name,row.amount_minor,row.source_version,row.source_time))
            if fail_batch==start:raise RuntimeError('injected before batch/checkpoint commit')
            conn.execute('UPDATE imports SET checkpoint=?,state=? WHERE tenant_id=? AND import_key=?',
                         (end,'complete' if end==len(records) else 'running',tenant_id,import_key))
        if after_commit:after_commit(end)

def spreadsheet_safe_cell(text):
    return "'"+text if text.lstrip().startswith(('=','+','-','@')) else text
